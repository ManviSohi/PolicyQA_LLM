import hashlib
from io import BytesIO
from uuid import UUID

from docx import Document as WordDocument
from fastapi.testclient import TestClient
from sqlalchemy import func, select
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

from app import database
from app.database import Base
from app.main import app
from app.models import Chunk, Document
from app.services.document_ingestion import MAX_FILE_SIZE


def test_pdf_upload_is_persisted_and_returned_by_document_endpoints(
    client: TestClient, pdf_factory,
) -> None:
    pdf_bytes = pdf_factory("Employee leave policy. " + "Policy details apply. " * 45)

    upload = client.post(
        "/api/documents",
        files={"file": ("leave-policy.pdf", pdf_bytes, "application/pdf")},
    )
    assert upload.status_code == 201, upload.text
    uploaded_document = upload.json()
    assert uploaded_document["status"] == "indexed"
    assert uploaded_document["chunk_count"] >= 1
    assert uploaded_document["sha256"] == hashlib.sha256(pdf_bytes).hexdigest()
    assert uploaded_document["created_at"]
    assert uploaded_document["updated_at"]

    document_list = client.get("/api/documents")
    assert document_list.status_code == 200
    assert len(document_list.json()) == 1
    assert document_list.json()[0]["id"] == uploaded_document["id"]

    detail = client.get(f"/api/documents/{uploaded_document['id']}")
    assert detail.status_code == 200
    assert detail.json()["status"] == "indexed"
    assert detail.json()["chunk_count"] == len(detail.json()["chunks"])
    assert detail.json()["chunks"][0]["document_id"] == uploaded_document["id"]
    assert detail.json()["chunks"][0]["page"] == 1

    duplicate = client.post(
        "/api/documents",
        files={"file": ("copy.pdf", pdf_bytes, "application/pdf")},
    )
    assert duplicate.status_code == 409
    with database.SessionLocal() as session:
        assert session.scalar(select(func.count()).select_from(Document)) == 1
        assert session.scalar(select(func.count()).select_from(Chunk)) == uploaded_document["chunk_count"]
        stored_document = session.get(Document, UUID(uploaded_document["id"]))
        stored_chunk = session.scalar(
            select(Chunk).where(Chunk.document_id == UUID(uploaded_document["id"]))
        )
        assert stored_document is not None
        assert stored_document.sha256 == uploaded_document["sha256"]
        assert stored_document.version == 1
        assert stored_document.status == "indexed"
        assert stored_chunk is not None
        assert stored_chunk.text.strip()
        assert stored_chunk.token_count > 0
        assert stored_chunk.page == 1


def test_docx_upload_preserves_heading_and_unknown_page(client: TestClient) -> None:
    document = WordDocument()
    document.add_heading("Leave Policy", level=1)
    document.add_paragraph("Employees may request leave through their manager.")
    content = BytesIO()
    document.save(content)

    response = client.post(
        "/api/documents",
        files={"file": ("leave-policy.docx", content.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )

    assert response.status_code == 201, response.text
    detail = client.get(f"/api/documents/{response.json()['id']}").json()
    assert detail["chunks"][0]["section"] == "Leave Policy"
    assert detail["chunks"][0]["page"] is None


def test_scanned_pdf_returns_clear_error_and_failed_status(client: TestClient, pdf_factory) -> None:
    response = client.post(
        "/api/documents",
        files={"file": ("scan.pdf", pdf_factory(""), "application/pdf")},
    )

    assert response.status_code == 422
    assert response.json()["detail"] == "Scanned PDF text extraction is not supported yet."
    assert client.get("/api/documents").json()[0]["status"] == "failed"


def test_invalid_uploads_return_useful_client_errors(client: TestClient, pdf_factory) -> None:
    unsupported = client.post(
        "/api/documents", files={"file": ("policy.txt", b"policy", "text/plain")}
    )
    assert unsupported.status_code == 400
    assert "PDF and DOCX" in unsupported.json()["detail"]

    empty = client.post("/api/documents", files={"file": ("empty.pdf", b"", "application/pdf")})
    assert empty.status_code == 400
    assert "empty" in empty.json()["detail"].lower()

    corrupted = client.post(
        "/api/documents", files={"file": ("corrupt.pdf", b"not a valid pdf", "application/pdf")}
    )
    assert corrupted.status_code == 422
    assert "extraction" in corrupted.json()["detail"].lower()
    assert client.get("/api/documents").json()[0]["status"] == "failed"

    oversized = client.post(
        "/api/documents",
        files={"file": ("large.pdf", b"x" * (MAX_FILE_SIZE + 1), "application/pdf")},
    )
    assert oversized.status_code == 413
    assert "25 MB" in oversized.json()["detail"]


def test_document_persists_across_api_restart(monkeypatch, tmp_path, pdf_factory) -> None:
    database_url = f"sqlite+pysqlite:///{(tmp_path / 'persistent.db').as_posix()}"
    first_engine = create_engine(database_url, connect_args={"check_same_thread": False})
    Base.metadata.create_all(bind=first_engine)
    monkeypatch.setattr(
        database,
        "SessionLocal",
        sessionmaker(bind=first_engine, autoflush=False, expire_on_commit=False),
    )
    monkeypatch.setattr(database, "create_schema", lambda: None)

    with TestClient(app) as first_client:
        uploaded = first_client.post(
            "/api/documents",
            files={"file": ("persistent.pdf", pdf_factory("Employees may request leave."), "application/pdf")},
        )
        assert uploaded.status_code == 201, uploaded.text
        document_id = uploaded.json()["id"]
    first_engine.dispose()

    restarted_engine = create_engine(database_url, connect_args={"check_same_thread": False})
    monkeypatch.setattr(
        database,
        "SessionLocal",
        sessionmaker(bind=restarted_engine, autoflush=False, expire_on_commit=False),
    )
    with TestClient(app) as restarted_client:
        documents = restarted_client.get("/api/documents").json()
        assert any(document["id"] == document_id for document in documents)
        restored = restarted_client.get(f"/api/documents/{document_id}").json()
        assert restored["status"] == "indexed"
        assert restored["chunk_count"] > 0
    restarted_engine.dispose()